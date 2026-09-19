"""
Ayush Knowledge Curator for Sanjeevani 2.0
Extracts, sanitizes, and categorizes:
1. vaidya_chikitsha.docx (114 Clinical Ailment Chapters) -> vaidya_chikitsa_curated.json
2. BotanicalHerb.docx (119 Medicinal Dravyaguna Herbs) -> botanical_herbs_curated.json

Enforces clinical safety:
- Divides ailments into 'household_safe', 'requires_consultation', and 'clinical_emergency'.
- Separates safe household lifestyle/herbal treatments from prescription-only classical formulations.
- Ingests Dravyaguna energetic profiles (hot/cold, Dosha actions) for intelligent advice.
"""

import os
import re
import json
import docx
from typing import List, Dict, Any, Tuple


# Ailment conditions strictly flagged as Emergency or requiring immediate hospital care.
# These will NEVER be recommended as self-care home remedies.
EMERGENCY_AILMENT_PATTERNS = [
    r"bone\s+fracture", r"convulsion", r"spasm", r"facial\s+paralysis",
    r"paralysis", r"tetanus", r"heart\s+disease", r"syncope", r"fainting",
    r"unconscious", r"pneumonia", r"cholera", r"plague", r"epilepsy",
    r"dengue\s+fever", r"bleeding\s+disorders", r"leprosy", r"chest\s+wound",
    r"menorrhagia", r"miscarriage", r"labor\s+pain", r"acute\s+nephritis",
    r"fistula", r"rectal\s+prolapse"
]

# Sub-acute ailments requiring medical/ASHA supervision, but safe for supportive Diet/Pathya
CONSULTATION_AILMENT_PATTERNS = [
    r"chickenpox", r"measles", r"rheumatism", r"typhoid", r"jaundice",
    r"chronic\s+fever", r"relapsing\s+fever", r"whooping\s+cough", r"asthma",
    r"diabetes", r"urinary\s+stones", r"kidney", r"liver\s+disease",
    r"goiter", r"ascites", r"pleurisy", r"gonorrhea", r"syphilis",
    r"deep\s+sinus", r"carbuncle"
]


def classify_safety_tier(title: str) -> str:
    """Classifies an ailment into household_safe, requires_consultation, or clinical_emergency."""
    t_lower = title.lower()
    for pat in EMERGENCY_AILMENT_PATTERNS:
        if re.search(pat, t_lower):
            return "clinical_emergency"
    for pat in CONSULTATION_AILMENT_PATTERNS:
        if re.search(pat, t_lower):
            return "requires_consultation"
    return "household_safe"


def extract_botanical_herbs(docx_path: str) -> List[Dict[str, Any]]:
    """Parses BotanicalHerb.docx (119 medicinal herbs with Dravyaguna properties)."""
    if not os.path.exists(docx_path):
        print(f"[Curator] Botanical herb file not found at: {docx_path}")
        return []

    doc = docx.Document(docx_path)
    herbs = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text or text.startswith("Botanical Herbs"):
            continue

        # Format: "Aadu (Ginger): Hot, increases digestive fire, destroys Vata and Kapha..."
        m = re.match(r"^([A-Za-z\s]+)\s*\(([^)]+)\):\s*(.*)", text)
        if m:
            vernacular = m.group(1).strip()
            common = m.group(2).strip()
            desc = m.group(3).strip()

            desc_lower = desc.lower()
            is_hot = "hot" in desc_lower or "pungent" in desc_lower
            is_cold = "cold" in desc_lower or "cooling" in desc_lower

            pacifies_vata = "vata" in desc_lower and ("destroys" in desc_lower or "pacifies" in desc_lower or "cures" in desc_lower)
            pacifies_pitta = "pitta" in desc_lower and ("destroys" in desc_lower or "pacifies" in desc_lower or "cures" in desc_lower)
            pacifies_kapha = "kapha" in desc_lower and ("destroys" in desc_lower or "pacifies" in desc_lower or "cures" in desc_lower)

            herbs.append({
                "herb_id": f"herb_{len(herbs)+1:03d}",
                "vernacular_name": vernacular,
                "common_name": common,
                "full_title": f"{vernacular} ({common})",
                "dravyaguna_properties": desc,
                "thermal_energy": "Hot (Ushna)" if is_hot else ("Cold (Sheetal)" if is_cold else "Neutral"),
                "dosha_balance": {
                    "pacifies_vata": pacifies_vata,
                    "pacifies_pitta": pacifies_pitta,
                    "pacifies_kapha": pacifies_kapha,
                },
                "searchable_summary": f"{vernacular} ({common}): {desc}"
            })

    return herbs


def extract_vaidya_chikitsa(docx_path: str) -> List[Dict[str, Any]]:
    """Parses vaidya_chikitsha.docx (114 clinical disease management chapters)."""
    if not os.path.exists(docx_path):
        print(f"[Curator] Vaidya Chikitsa file not found at: {docx_path}")
        return []

    doc = docx.Document(docx_path)
    ailments = []
    current = None

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        # Match chapter title: e.g. "1. Indigestion (Ajirna / Apacho)"
        m = re.match(r"^(\d+)\.\s+(.*)", text)
        if m:
            if current:
                ailments.append(current)
            
            title = m.group(2).strip()
            tier = classify_safety_tier(title)

            current = {
                "id": f"vc_{int(m.group(1)):03d}",
                "ailment_index": int(m.group(1)),
                "title": title,
                "safety_tier": tier,
                "causes": "",
                "symptoms": "",
                "general_treatment": "",
                "classical_medicines": "",
                "diet_pathya": "",
                "current_field": "intro"
            }
            continue

        if current:
            if text.startswith("Causes:") or text.startswith("Causes & Symptoms:"):
                current["current_field"] = "causes"
                current["causes"] += text.split(":", 1)[1].strip() + " "
            elif text.startswith("Symptoms:"):
                current["current_field"] = "symptoms"
                current["symptoms"] += text.split(":", 1)[1].strip() + " "
            elif text.startswith("General Treatment:"):
                current["current_field"] = "general_treatment"
                current["general_treatment"] += text.split(":", 1)[1].strip() + " "
            elif text.startswith("Classical Ayurvedic Medicines:"):
                current["current_field"] = "classical_medicines"
                current["classical_medicines"] += text.split(":", 1)[1].strip() + " "
            elif text.startswith("Diet (Pathya):") or text.startswith("Diet:"):
                current["current_field"] = "diet_pathya"
                current["diet_pathya"] += text.split(":", 1)[1].strip() + " "
            else:
                f = current["current_field"]
                if f in current and f != "intro":
                    current[f] += text + " "

    if current:
        ailments.append(current)

    # Clean and enrich entries for RAG
    curated_ailments = []
    for item in ailments:
        # Build search-focused index string (prevents vector smearing)
        title_clean = item["title"].strip()
        symptoms_clean = item["symptoms"].strip() or item["causes"].strip()
        treatment_clean = item["general_treatment"].strip()
        pathya_clean = item["diet_pathya"].strip()
        meds_clean = item["classical_medicines"].strip()

        # Format as standard Sanjeevani remedy entry
        entry = {
            "remedy_id": item["id"],
            "remedy_name": title_clean,
            "category": "Classical Vaidya Chikitsa",
            "safety_tier": item["safety_tier"],
            "indications": title_clean,
            "symptoms": [s.strip() for s in re.split(r"[,;.]", symptoms_clean) if len(s.strip()) > 3][:6],
            "causes": item["causes"].strip(),
            "preparation": treatment_clean or "Follow lifestyle and dietary fasting (Langhan).",
            "diet_pathya": pathya_clean or "Simple, warm, light digestible food (Khichdi, warm water).",
            "dosage": "As recommended under general treatment or physician guidance.",
            "classical_medicines_note": meds_clean or "None",
            "requires_doctor_visit": item["safety_tier"] != "household_safe",
            "safety_precaution": (
                "For severe symptoms or persistence > 2 days, consult your primary health center or ASHA worker immediately."
                if item["safety_tier"] == "household_safe"
                else "Requires evaluation by an Ayurvedic Medical Officer or Primary Health Center."
            ),
            # Key for vector search: Only clinical indications & symptoms!
            "searchable_text": f"Condition: {title_clean}. Symptoms: {symptoms_clean[:200]}."
        }
        curated_ailments.append(entry)

    return curated_ailments


def build_and_save_curated_knowledge(data_dir: str = "DATA/Ayush", output_dir: str = "DATA"):
    """Runs extraction on both docx files and saves standardized JSON datasets."""
    # Handle different execution CWDs
    if not os.path.exists(data_dir) and os.path.exists(os.path.join("backend", data_dir)):
        data_dir = os.path.join("backend", data_dir)
    if not os.path.exists(output_dir) and os.path.exists(os.path.join("backend", output_dir)):
        output_dir = os.path.join("backend", output_dir)

    os.makedirs(output_dir, exist_ok=True)

    botanical_path = os.path.join(data_dir, "BotanicalHerb.docx")
    vaidya_path = os.path.join(data_dir, "vaidya_chikitsha.docx")

    print(f"[Knowledge Curator] Processing {botanical_path}...")
    herbs = extract_botanical_herbs(botanical_path)
    herbs_out = os.path.join(output_dir, "botanical_herbs_curated.json")
    with open(herbs_out, "w", encoding="utf-8") as f:
        json.dump(herbs, f, indent=2, ensure_ascii=False)
    print(f"[Knowledge Curator] Saved {len(herbs)} botanical herbs to {herbs_out}")

    print(f"[Knowledge Curator] Processing {vaidya_path}...")
    ailments = extract_vaidya_chikitsa(vaidya_path)
    ailments_out = os.path.join(output_dir, "vaidya_chikitsa_curated.json")
    with open(ailments_out, "w", encoding="utf-8") as f:
        json.dump(ailments, f, indent=2, ensure_ascii=False)
    print(f"[Knowledge Curator] Saved {len(ailments)} clinical ailments to {ailments_out}")

    safe_count = sum(1 for a in ailments if a["safety_tier"] == "household_safe")
    consult_count = sum(1 for a in ailments if a["safety_tier"] == "requires_consultation")
    emerg_count = sum(1 for a in ailments if a["safety_tier"] == "clinical_emergency")

    print(f"[Knowledge Curator] Tiers: {safe_count} Household-Safe, {consult_count} Require Consultation, {emerg_count} Clinical Emergency.")
    return len(herbs), len(ailments)


if __name__ == "__main__":
    build_and_save_curated_knowledge()
